"""
Generate a raw UNFORMATTED sample_manuscript.docx using Book 53 content.
The document has deliberately inconsistent formatting (random fonts, sizes,
alignments) so the PS 1 formatter tool can process and reformat it.

Content includes: title page, author, TOC, 15 chapters with subheadings,
tables (numbered), bullet points, MCQ blocks, figure placeholders, index,
and references — all with structural markers the rule engine can detect.
"""

import os
import random
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import docx

random.seed(42)

FONTS = ['Arial', 'Calibri', 'Times New Roman', 'Courier New', 'Verdana', 'Georgia', 'Tahoma']
SIZES = [Pt(10), Pt(11), Pt(12), Pt(13), Pt(14), Pt(15)]
ALIGNMENTS = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY]
SPACINGS = [1.0, 1.15, 1.5, 2.0]

INPUT_FILE = r"Book no 53 Self-Evolving Enterprise Systems_ AI-Driven Architectures for the Next Generation of Digital Platforms.docx"
OUTPUT_FILE = r"sample_manuscript.docx"

BOOK_TITLE = "Self-Evolving Enterprise Systems: AI-Driven Architectures for the Next Generation of Digital Platforms"
AUTHOR_NAME = "Author Name"


def add_raw_paragraph(doc, text, is_heading=False):
    """Add a paragraph with deliberately inconsistent formatting."""
    p = doc.add_paragraph()
    run = p.add_run(text)

    font_name = random.choice(FONTS)
    font_size = random.choice(SIZES)
    alignment = random.choice(ALIGNMENTS)
    spacing = random.choice(SPACINGS)

    if is_heading:
        run.bold = random.choice([True, False])
        if random.random() > 0.5:
            font_size = Pt(random.randint(14, 20))
        else:
            font_size = Pt(random.randint(10, 13))
        if random.random() > 0.7:
            run.text = run.text.upper()
    else:
        run.bold = False

    p.alignment = alignment
    if p.paragraph_format:
        p.paragraph_format.line_spacing = spacing
        if random.random() > 0.5:
            p.paragraph_format.first_line_indent = Cm(random.choice([0, 0.5, 1.0, 1.27, 2.0]))

    run.font.name = font_name
    run.font.size = font_size
    return p


def add_raw_table(doc, headers, rows):
    """Add a table with no consistent formatting."""
    num_cols = len(headers)
    num_rows = len(rows) + 1
    table = doc.add_table(rows=num_rows, cols=num_cols)

    for ci, header in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = header
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.name = random.choice(FONTS)
                run.font.size = Pt(random.randint(9, 13))
                run.bold = random.choice([True, False])

    for ri, row_data in enumerate(rows):
        for ci, cell_text in enumerate(row_data):
            cell = table.cell(ri + 1, ci)
            cell.text = str(cell_text)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = random.choice(FONTS)
                    run.font.size = Pt(random.randint(9, 12))

    return table


def extract_chapters(input_file):
    """Extract chapter content from Book 53."""
    src_doc = Document(input_file)
    chapters = {}
    chapter_order = []
    current_chapter = None
    current_paras = []

    for p in src_doc.paragraphs:
        if p.style and p.style.name == 'Heading 1' and p.text.strip():
            if current_chapter and current_paras:
                chapters[current_chapter] = current_paras
            title = p.text.strip()
            if title != 'Table of Contents':
                current_chapter = title
                chapter_order.append(title)
                current_paras = []
            else:
                current_chapter = None
                current_paras = []
        elif current_chapter and p.text.strip():
            current_paras.append(p.text.strip())

    if current_chapter and current_paras:
        chapters[current_chapter] = current_paras

    return chapters, chapter_order


# ============================================================
# SUBHEADINGS PER CHAPTER
# ============================================================
SUBHEADINGS = {
    1: ["The Historical Context of Enterprise Evolution",
        "Understanding Self-Evolution in Digital Platforms",
        "The Convergence of Cloud, AI, and Big Data",
        "Biological Metaphors for Software Architecture",
        "The Role of Autonomous Decision-Making",
        "Challenges in Transitioning to Self-Evolving Systems",
        "The Human Element in Autonomous Platforms",
        "Key Principles of Self-Evolution",
        "Summary and Key Takeaways"],
    2: ["The Rise and Fall of Monolithic Architecture",
        "The Microservices Revolution",
        "Service Mesh and Infrastructure Intelligence",
        "Event-Driven Architecture in the Mesh",
        "Domain-Driven Design and Bounded Contexts",
        "Container Orchestration and Kubernetes",
        "The Strangler Fig Migration Pattern",
        "API Gateway and Traffic Management",
        "Observability in Distributed Systems",
        "Summary and Key Takeaways"],
    3: ["Machine Learning Paradigms Overview",
        "Neural Networks and Deep Learning Foundations",
        "Hardware Acceleration for AI Workloads",
        "Transformer Architecture and Large Language Models",
        "Vector Databases and Semantic Search",
        "Transfer Learning and Pre-trained Models",
        "Real-Time Data Streaming for AI Systems",
        "Federated Learning and Privacy-Preserving AI",
        "The MAPE-K Autonomous Management Loop",
        "MLOps and Model Lifecycle Management",
        "Summary and Key Takeaways"],
    4: ["Understanding Cognitive Load in AI Systems",
        "Hardware Orchestration and Dynamic Allocation",
        "Token Economics for Large Language Models",
        "Hierarchical Model Orchestration",
        "Semantic Memory Management",
        "Predictive Scaling with Causal AI",
        "FinOps for AI-Driven Enterprises",
        "Green AI and Energy Efficiency",
        "Edge Computing Resource Management",
        "Summary and Key Takeaways"],
    5: ["Principles of Self-Healing Architecture",
        "Circuit Breaker and Bulkhead Patterns",
        "Autonomous Failure Detection and Recovery",
        "Chaos Engineering for Resilience Testing",
        "Distributed Consensus and Consistency",
        "Data Integrity and Automatic Repair",
        "Anti-Fragile System Design",
        "Graceful Degradation Strategies",
        "Monitoring and Alerting for Resilience",
        "Summary and Key Takeaways"],
    6: ["The Promise of Generative AI in Enterprise Development",
        "AI-Assisted Database Schema Design",
        "Automated Code Synthesis and Review",
        "Test Generation and Quality Assurance",
        "Natural Language to API Translation",
        "Model-Driven Architecture with AI",
        "Risks and Limitations of Generated Code",
        "Human-AI Collaboration in Development",
        "Summary and Key Takeaways"],
    7: ["Defining the Intelligent Data Fabric",
        "Data Catalog and Metadata Management",
        "Automated Data Quality Frameworks",
        "Data Governance and Compliance",
        "Real-Time vs. Batch Processing Architectures",
        "Data Virtualization and Unified Access",
        "Knowledge Graphs in Data Integration",
        "Data Lineage and Impact Analysis",
        "Summary and Key Takeaways"],
    8: ["Foundations of Reinforcement Learning",
        "Reward Function Design for Business Objectives",
        "Multi-Agent Reinforcement Learning",
        "Exploration vs. Exploitation in Enterprise Context",
        "Applying RL to Infrastructure Optimization",
        "Continuous Deployment with RL Feedback",
        "Safe Reinforcement Learning Practices",
        "Case Studies in RL for Enterprise Improvement",
        "Summary and Key Takeaways"],
    9: ["The Evolving Threat Landscape",
        "Zero Trust Architecture Principles",
        "AI-Powered Threat Detection",
        "Autonomous Incident Response",
        "Identity and Access Management with AI",
        "Data Protection and Encryption Strategies",
        "Supply Chain Security for AI Models",
        "Security Operations Center Automation",
        "Regulatory Compliance and Audit",
        "Summary and Key Takeaways"],
    10: ["The Need for Explainable AI",
         "Post-Hoc Explainability Methods",
         "Intrinsic Interpretability Approaches",
         "Building Observable AI Systems",
         "Distributed Tracing for AI Pipelines",
         "Metrics, Logs, and Traces Integration",
         "AI Model Monitoring in Production",
         "Audit Trails and Compliance Logging",
         "Summary and Key Takeaways"],
    11: ["Horizontal and Vertical Scaling Strategies",
         "Cloud-Native Architecture for Scale",
         "Database Scaling and Sharding",
         "Global Distribution and Data Replication",
         "AI-Driven Auto-Scaling Policies",
         "Cost-Efficient Scaling Patterns",
         "Performance Testing at Scale",
         "Multi-Cloud and Hybrid Cloud Scaling",
         "Summary and Key Takeaways"],
    12: ["The Evolution from DevOps to AIOps",
         "Automated Incident Management",
         "AI-Driven Root Cause Analysis",
         "Intelligent Change Management",
         "Predictive Operations and Capacity Planning",
         "AI-Enhanced CI/CD Pipelines",
         "Service Level Objectives and AI",
         "The Future of Software Lifecycle Automation",
         "Summary and Key Takeaways"],
    13: ["Ethical Frameworks for AI Governance",
         "Bias Detection and Mitigation",
         "Transparency and Accountability Structures",
         "Privacy-Preserving AI Techniques",
         "Human Agency in Autonomous Systems",
         "Regulatory Landscape for AI",
         "Organizational AI Governance Structures",
         "Building Responsible AI Teams",
         "Summary and Key Takeaways"],
    14: ["Planning for Self-Evolving System Implementation",
         "Technology Selection and Evaluation",
         "Team Structure and Skills Development",
         "Pilot Program Design and Execution",
         "Case Study: Financial Services Transformation",
         "Case Study: Healthcare AI Integration",
         "Case Study: Manufacturing Intelligence",
         "Lessons Learned and Best Practices",
         "Summary and Key Takeaways"],
    15: ["Emerging Technologies on the Horizon",
         "Quantum Computing and Enterprise AI",
         "Neuromorphic Computing for Edge Intelligence",
         "The Path Toward Artificial General Intelligence",
         "Autonomous Architecture Generation",
         "The Role of Digital Twins",
         "Sustainability and Green Enterprise Computing",
         "Preparing Organizations for the Future",
         "Summary and Key Takeaways"],
}

# ============================================================
# TABLES PER CHAPTER
# ============================================================
TABLES = {
    1: [("Evolution of Enterprise Systems",
         ["Era", "Architecture", "Key Characteristic", "Limitation"],
         [["1960s-1980s", "Mainframe", "Centralized processing", "Limited accessibility"],
          ["1980s-2000s", "Client-Server", "Distributed computing", "Complex maintenance"],
          ["2000s-2015", "Service-Oriented", "Reusable services", "Integration overhead"],
          ["2015-2020", "Microservices", "Independent deployment", "Orchestration complexity"],
          ["2020-Present", "Self-Evolving", "Autonomous adaptation", "Governance challenges"]]),
        ("Key Drivers of Self-Evolving Enterprise Systems",
         ["Driver", "Description", "Impact"],
         [["Data Explosion", "Exponential growth in data volumes", "Requires automated processing"],
          ["Cloud Computing", "On-demand scalable infrastructure", "Enables elastic architectures"],
          ["AI/ML Advances", "Sophisticated learning algorithms", "Powers autonomous decision-making"],
          ["Global Competition", "Rapid market changes", "Demands agile adaptation"],
          ["Digital Transformation", "Business process digitization", "Creates integration complexity"]])],
    2: [("Monolithic vs. Microservices Comparison",
         ["Aspect", "Monolithic", "Microservices", "Intelligent Mesh"],
         [["Deployment", "Single unit", "Independent services", "AI-orchestrated services"],
          ["Scaling", "Vertical only", "Horizontal per service", "Autonomous scaling"],
          ["Technology", "Single stack", "Polyglot", "AI-selected optimal stack"],
          ["Resilience", "Single point of failure", "Isolated failures", "Self-healing"],
          ["Development", "Tightly coupled teams", "Independent teams", "AI-assisted development"]]),
        ("Service Mesh Technologies",
         ["Technology", "Purpose", "Key Feature"],
         [["Istio", "Service mesh platform", "Traffic management and security"],
          ["Envoy", "Edge and service proxy", "High-performance data plane"],
          ["Linkerd", "Lightweight service mesh", "Simplicity and performance"],
          ["Consul", "Service discovery", "Multi-datacenter support"],
          ["NGINX", "Reverse proxy", "Load balancing and caching"]])],
    3: [("AI/ML Framework Comparison",
         ["Framework", "Language", "Strengths", "Use Case"],
         [["TensorFlow", "Python/C++", "Production deployment", "Deep learning at scale"],
          ["PyTorch", "Python", "Research flexibility", "Rapid prototyping"],
          ["scikit-learn", "Python", "Classical ML algorithms", "Data analysis"],
          ["Spark MLlib", "Scala/Python", "Distributed processing", "Big data ML"],
          ["Hugging Face", "Python", "NLP models", "Language understanding"]]),
        ("Types of Neural Network Architectures",
         ["Architecture", "Description", "Application"],
         [["CNN", "Convolutional Neural Network", "Image recognition"],
          ["RNN", "Recurrent Neural Network", "Sequential data"],
          ["Transformer", "Attention-based architecture", "NLP and generative AI"],
          ["GAN", "Generative Adversarial Network", "Synthetic data"],
          ["Autoencoder", "Compression and reconstruction", "Anomaly detection"]])],
    4: [("Resource Types in Cognitive Systems",
         ["Resource", "Unit", "Management Strategy", "Optimization"],
         [["Compute", "vCPUs/CUDA cores", "Dynamic allocation", "Predictive scaling"],
          ["Memory", "GB/TB", "Tiered caching", "Intelligent prefetching"],
          ["Storage", "TB/PB", "Data gravity management", "Automated tiering"],
          ["Tokens", "Token count", "Token economics", "Model routing"],
          ["Network", "Gbps", "Traffic shaping", "Edge distribution"]]),
        ("Cost Optimization Strategies",
         ["Strategy", "Savings Potential", "Trade-off"],
         [["Model distillation", "40-60%", "Slight accuracy reduction"],
          ["Spot instances", "60-80%", "Interruption risk"],
          ["Dynamic model switching", "30-50%", "Latency variation"],
          ["Prompt caching", "20-40%", "Memory overhead"],
          ["Edge inference", "25-45%", "Limited model size"]])],
    5: [("Self-Healing Mechanisms",
         ["Mechanism", "Detection Method", "Recovery Action", "Time to Recovery"],
         [["Circuit breaker", "Failure rate threshold", "Isolate service", "Milliseconds"],
          ["Auto-restart", "Health check failure", "Restart container", "Seconds"],
          ["Failover", "Primary node failure", "Switch to replica", "Sub-second"],
          ["Auto-scaling", "Load threshold breach", "Spawn new instances", "Minutes"],
          ["Data repair", "Corruption detection", "Restore from backup", "Minutes to hours"]]),
        ("Resilience Patterns",
         ["Pattern", "Purpose", "Implementation"],
         [["Bulkhead", "Isolate failures", "Separate thread pools per service"],
          ["Retry with backoff", "Handle transient errors", "Exponential delay"],
          ["Timeout", "Prevent resource starvation", "Maximum wait times"],
          ["Fallback", "Graceful degradation", "Default response on failure"],
          ["Chaos engineering", "Proactive testing", "Controlled failure injection"]])],
    6: [("Generative AI Applications",
         ["Application", "Input", "Output", "Benefit"],
         [["Schema generation", "Business requirements", "Database schemas", "Faster development"],
          ["Code synthesis", "Natural language specs", "Working code", "Reduced dev time"],
          ["API design", "Service descriptions", "OpenAPI specs", "Consistent interfaces"],
          ["Test generation", "Source code", "Unit tests", "Better coverage"],
          ["Documentation", "Code and configs", "Technical docs", "Up-to-date docs"]])],
    7: [("Data Fabric Components",
         ["Component", "Function", "Technology Example"],
         [["Data Catalog", "Metadata management", "Apache Atlas, Alation"],
          ["Data Integration", "ETL/ELT processing", "Apache NiFi, Talend"],
          ["Data Quality", "Validation and cleansing", "Great Expectations"],
          ["Data Governance", "Policy enforcement", "Apache Ranger"],
          ["Data Virtualization", "Unified access layer", "Denodo, Dremio"]])],
    8: [("RL Algorithms for Enterprise",
         ["Algorithm", "Type", "Enterprise Application"],
         [["Q-Learning", "Model-free, off-policy", "Resource allocation"],
          ["SARSA", "Model-free, on-policy", "Risk-sensitive operations"],
          ["DQN", "Deep RL", "Infrastructure optimization"],
          ["PPO", "Policy gradient", "Workflow orchestration"],
          ["A3C", "Actor-critic", "Multi-agent coordination"]])],
    9: [("Security Layers",
         ["Layer", "Protection Scope", "AI Technique", "Response Time"],
         [["Perimeter", "Network boundary", "Anomaly detection", "Real-time"],
          ["Application", "Service logic", "Behavioral analysis", "Near-real-time"],
          ["Data", "Information assets", "Encryption and DLP", "Continuous"],
          ["Identity", "User and service auth", "Zero-trust verification", "Per-request"],
          ["Infrastructure", "Cloud and hardware", "Drift detection", "Periodic scan"]])],
    10: [("Explainability Methods",
          ["Method", "Type", "Scope", "Best For"],
          [["LIME", "Post-hoc", "Local", "Individual prediction explanation"],
           ["SHAP", "Post-hoc", "Local/Global", "Feature importance analysis"],
           ["Attention vis", "Intrinsic", "Local", "Transformer models"],
           ["Counterfactual", "Post-hoc", "Local", "Decision boundary analysis"],
           ["Rule extraction", "Post-hoc", "Global", "Policy compliance"]])],
    11: [("Scaling Strategies",
          ["Strategy", "Direction", "AI Enhancement", "Challenge"],
          [["Vertical", "Up (bigger machines)", "Workload prediction", "Hardware limits"],
           ["Horizontal", "Out (more instances)", "Auto-scaling policies", "State management"],
           ["Functional", "By service type", "Intelligent routing", "Service discovery"],
           ["Geographic", "Across regions", "Data gravity aware", "Consistency"],
           ["Temporal", "By time patterns", "Predictive scheduling", "Accuracy"]])],
    12: [("AIOps Maturity Model",
          ["Level", "Stage", "Capabilities", "AI Integration"],
          [["1", "Reactive", "Manual monitoring", "None"],
           ["2", "Proactive", "Alert correlation", "Rule-based"],
           ["3", "Predictive", "Anomaly detection", "ML-driven"],
           ["4", "Autonomous", "Self-healing", "AI-native"],
           ["5", "Cognitive", "Continuous learning", "Fully autonomous"]])],
    13: [("AI Ethics Framework",
          ["Component", "Focus Area", "Implementation"],
          [["Transparency", "Decision explainability", "XAI methods and audit logs"],
           ["Fairness", "Bias prevention", "Fairness metrics and testing"],
           ["Accountability", "Responsibility assignment", "Governance structures"],
           ["Privacy", "Data protection", "Differential privacy, encryption"],
           ["Safety", "Harm prevention", "Guardrails and kill switches"]])],
    14: [("Industry Implementation Examples",
          ["Industry", "Use Case", "AI Technique", "Business Impact"],
          [["Finance", "Fraud detection", "Anomaly detection", "80% faster detection"],
           ["Healthcare", "Diagnostic support", "Computer vision", "95% accuracy improvement"],
           ["Manufacturing", "Predictive maintenance", "Time-series analysis", "40% downtime reduction"],
           ["Retail", "Demand forecasting", "Deep learning", "30% inventory optimization"],
           ["Energy", "Grid optimization", "Reinforcement learning", "25% efficiency gain"]])],
    15: [("Emerging Technologies",
          ["Technology", "Maturity", "Impact Potential", "Timeline"],
          [["Quantum computing", "Early research", "Revolutionary", "5-10 years"],
           ["Neuromorphic chips", "Prototype", "High", "3-7 years"],
           ["AGI", "Theoretical", "Transformative", "10+ years"],
           ["Brain-computer interfaces", "Early research", "Medium", "7-15 years"],
           ["Autonomous code gen", "Emerging", "Very high", "2-5 years"]])],
}

# ============================================================
# MCQs PER CHAPTER (15 per chapter)
# ============================================================
def generate_mcqs(chapter_num, chapter_title):
    """Generate 15 MCQs per chapter."""
    mcq_templates = [
        ("What is the primary goal of %s?",
         ["Increasing manual work", "Enabling autonomous, intelligent operations",
          "Reducing system capabilities", "Eliminating all technology"], "b"),
        ("Which technology is most critical for %s?",
         ["Manual spreadsheets", "AI and machine learning algorithms",
          "Paper-based tracking", "Fax machines"], "b"),
        ("What challenge is addressed by %s?",
         ["Reducing innovation", "Managing complexity in modern enterprise systems",
          "Simplifying to single-server", "Eliminating data collection"], "b"),
        ("How does %s improve enterprise operations?",
         ["By adding manual processes", "Through continuous autonomous optimization",
          "By slowing processing", "By removing automation"], "b"),
        ("What role does data play in %s?",
         ["No significant role", "Serves as the foundation for learning and adaptation",
          "Only for compliance reporting", "Data is irrelevant"], "b"),
        ("What is a key benefit of %s?",
         ["Increased downtime", "Enhanced resilience and faster response times",
          "More complex manual processes", "Higher operational costs"], "b"),
        ("Which pattern supports %s?",
         ["Monolithic architecture", "Distributed microservices with AI orchestration",
          "Single-server deployment", "Batch-only processing"], "b"),
        ("What is the relationship between observability and %s?",
         ["No relationship", "Observability provides data for autonomous decisions",
          "Observability replaces it", "They are mutually exclusive"], "b"),
        ("How does %s handle failure scenarios?",
         ["Ignores failures", "Implements self-healing and automated recovery",
          "Requires manual restart", "Shuts down permanently"], "b"),
        ("What ethical consideration is important for %s?",
         ["Ethics are not relevant", "Ensuring transparency, fairness, and human oversight",
          "Maximizing profit at all costs", "Eliminating human involvement"], "b"),
        ("What is the impact of %s on scalability?",
         ["Reduces scalability", "Enables dynamic, intelligent scaling based on demand",
          "Requires fixed infrastructure", "Eliminates scaling need"], "b"),
        ("How does %s integrate with existing systems?",
         ["Requires complete replacement", "Through gradual adoption and interoperability",
          "Cannot integrate with legacy", "Only works with new systems"], "b"),
        ("What metric measures success in %s?",
         ["Lines of code", "System availability, performance, and adaptation speed",
          "Number of servers", "Manual intervention frequency"], "b"),
        ("What is the future outlook for %s?",
         ["Becoming obsolete", "Increasing adoption with more sophisticated AI",
          "Remaining unchanged", "Being replaced by manual processes"], "b"),
        ("Which stakeholder benefits most from %s?",
         ["Only IT department", "The entire organization including business and technical teams",
          "Only external customers", "Only management"], "b"),
    ]
    topic = chapter_title.lower()
    mcqs = []
    for q_template, opts, answer in mcq_templates:
        mcqs.append({
            "question": q_template % topic,
            "options": opts,
            "answer": answer
        })
    return mcqs


# ============================================================
# FIGURE CAPTIONS PER CHAPTER
# ============================================================
FIGURE_CAPTIONS = {
    1: ["Evolution timeline of enterprise architectures from mainframe to self-evolving systems",
        "The self-evolving enterprise feedback loop: Monitor, Analyze, Adapt, Optimize",
        "Comparison of traditional vs. self-evolving system architectures"],
    2: ["Monolithic architecture vs. microservices vs. intelligent mesh topology",
        "Service mesh architecture with sidecar proxy pattern",
        "Event-driven communication flow in an intelligent mesh"],
    3: ["Neural network architecture layers: input, hidden, and output",
        "Transformer architecture with multi-head attention mechanism",
        "MLOps pipeline: from data ingestion to model deployment"],
    4: ["Cognitive resource management dashboard showing real-time allocation",
        "Hierarchical model orchestration: routing requests to optimal AI models",
        "Token economics flow diagram for LLM-based enterprise applications"],
    5: ["Self-healing architecture with circuit breaker pattern",
        "Chaos engineering experiment workflow",
        "Anti-fragile system response curve"],
    6: ["AI-assisted code generation pipeline",
        "Generative architecture workflow: schema to API",
        "Human-AI collaboration model in development"],
    7: ["Intelligent data fabric architecture",
        "Data pipeline comparison: batch vs. stream vs. lambda",
        "Knowledge graph of enterprise data relationships"],
    8: ["Reinforcement learning agent-environment loop",
        "Multi-agent RL coordination for resource optimization",
        "Reward function evolution over training episodes"],
    9: ["Zero trust security architecture",
        "AI-powered SOC workflow",
        "Threat detection and autonomous response timeline"],
    10: ["SHAP value visualization for AI model decision explanation",
         "Distributed tracing view of AI inference pipeline",
         "Observability stack integration"],
    11: ["Auto-scaling architecture with AI prediction engine",
         "Multi-region deployment topology",
         "Database sharding strategy visualization"],
    12: ["AIOps maturity model progression",
         "AI-enhanced CI/CD pipeline with quality gates",
         "Intelligent incident management workflow"],
    13: ["AI ethics framework visualization",
         "Human-AI collaboration spectrum",
         "Organizational AI governance structure"],
    14: ["Implementation roadmap phases",
         "Case study results: before and after deployment",
         "Key success factors for AI transformation"],
    15: ["Future technology convergence: quantum, neuromorphic, AI",
         "Digital twin architecture for enterprise simulation",
         "Vision of the fully autonomous enterprise platform"],
}

# ============================================================
# KEY POINTS (BULLET LISTS) PER CHAPTER
# ============================================================
BULLET_POINTS = {
    1: ["Self-evolving systems autonomously monitor, adapt, and optimize enterprise operations",
        "The convergence of cloud, AI, and big data enables autonomous digital platforms",
        "Biological evolution provides a powerful metaphor for understanding system adaptation",
        "The transition from static to dynamic systems is the defining challenge of modern IT",
        "Human oversight remains critical for governance, ethics, and strategic direction",
        "Generative AI accelerates self-evolution through code synthesis and configuration",
        "Observability and telemetry data form the foundation of autonomous decision-making",
        "New organizational roles and skills are required to manage self-evolving platforms"],
    2: ["Monolithic architectures served well initially but could not scale to modern demands",
        "Microservices enable independent deployment and scaling of individual components",
        "Service mesh provides infrastructure-level intelligence for service communication",
        "Event-driven architecture enables loose coupling and real-time responsiveness",
        "Domain-driven design aligns software architecture with business domains",
        "Kubernetes orchestrates containerized workloads across distributed infrastructure",
        "The strangler fig pattern enables gradual migration from monolith to microservices",
        "API gateways manage routing, security, and rate limiting for distributed services"],
}

# ============================================================
# INDEX TERMS
# ============================================================
INDEX_TERMS = [
    "Artificial Intelligence (AI)", "Machine Learning (ML)", "Deep Learning",
    "Neural Networks", "Transformer Architecture", "Microservices",
    "Service Mesh", "Kubernetes", "Docker", "Containerization",
    "Self-Healing Systems", "Circuit Breaker Pattern", "Bulkhead Pattern",
    "Chaos Engineering", "Reinforcement Learning", "Q-Learning",
    "Policy Gradient", "Actor-Critic", "Generative AI", "Large Language Models",
    "Natural Language Processing", "Computer Vision", "CNN", "RNN",
    "Vector Database", "Knowledge Graph", "Data Fabric", "Data Mesh",
    "Event-Driven Architecture", "Domain-Driven Design",
    "API Gateway", "GraphQL", "REST API",
    "Zero Trust Architecture", "IAM", "SIEM", "SOAR", "Threat Detection",
    "Anomaly Detection", "Predictive Analytics",
    "MLOps", "DevOps", "AIOps", "CI/CD Pipeline",
    "Auto-Scaling", "Load Balancing", "Edge Computing",
    "Cloud Computing", "AWS", "Azure", "Google Cloud Platform",
    "Serverless Computing", "Federated Learning", "Transfer Learning",
    "Token Economics", "Prompt Engineering", "RAG",
    "Embeddings", "Feature Store", "Model Registry",
    "Data Governance", "Data Quality", "Data Lineage",
    "Observability", "Distributed Tracing",
    "Explainable AI", "LIME", "SHAP",
    "Bias Detection", "Fairness in AI", "Ethics in AI",
    "GDPR", "CCPA", "Regulatory Compliance",
    "Digital Twin", "Quantum Computing", "Neuromorphic Computing",
    "FinOps", "Green AI", "Sustainability",
    "Technical Debt", "Scalability", "High Availability",
    "Database Sharding", "Replication",
    "Apache Kafka", "Apache Spark", "TensorFlow", "PyTorch",
    "Autonomous Systems", "Self-Evolving Architecture",
    "MAPE-K Loop", "Cognitive Computing",
    "Model Drift", "Data Drift",
    "Canary Deployment", "Blue-Green Deployment",
    "Infrastructure as Code", "Terraform",
    "Agile Methodology", "Site Reliability Engineering",
]

# ============================================================
# ABBREVIATIONS
# ============================================================
ABBREVIATIONS = [
    ("AI", "Artificial Intelligence"), ("ML", "Machine Learning"),
    ("API", "Application Programming Interface"), ("AWS", "Amazon Web Services"),
    ("CI/CD", "Continuous Integration / Continuous Deployment"),
    ("CNN", "Convolutional Neural Network"), ("CPU", "Central Processing Unit"),
    ("DDD", "Domain-Driven Design"), ("DevOps", "Development and Operations"),
    ("GPU", "Graphics Processing Unit"), ("IAM", "Identity and Access Management"),
    ("IoT", "Internet of Things"), ("LLM", "Large Language Model"),
    ("MAPE-K", "Monitor, Analyze, Plan, Execute - Knowledge"),
    ("MLOps", "Machine Learning Operations"), ("NLP", "Natural Language Processing"),
    ("RL", "Reinforcement Learning"), ("RNN", "Recurrent Neural Network"),
    ("SIEM", "Security Information and Event Management"),
    ("SLA", "Service Level Agreement"), ("SLO", "Service Level Objective"),
    ("SOC", "Security Operations Center"), ("SRE", "Site Reliability Engineering"),
    ("TPU", "Tensor Processing Unit"), ("XAI", "Explainable AI"),
]

# ============================================================
# REFERENCES
# ============================================================
REFERENCES = [
    "[1] Russell, S. and Norvig, P. (2021). Artificial Intelligence: A Modern Approach, 4th Edition. Pearson.",
    "[2] Goodfellow, I., Bengio, Y., and Courville, A. (2016). Deep Learning. MIT Press.",
    "[3] Kleppmann, M. (2017). Designing Data-Intensive Applications. O'Reilly Media.",
    "[4] Newman, S. (2021). Building Microservices, 2nd Edition. O'Reilly Media.",
    "[5] Sutton, R. and Barto, A. (2018). Reinforcement Learning: An Introduction. MIT Press.",
    "[6] Vaswani, A. et al. (2017). Attention Is All You Need. NeurIPS, 30.",
    "[7] Burns, B. et al. (2016). Borg, Omega, and Kubernetes. ACM Queue, 14(1).",
    "[8] Brown, T. et al. (2020). Language Models are Few-Shot Learners. NeurIPS, 33.",
    "[9] Nygard, M. (2018). Release It! Design and Deploy Production-Ready Software. Pragmatic Bookshelf.",
    "[10] Huyen, C. (2022). Designing Machine Learning Systems. O'Reilly Media.",
    "[11] Kephart, J. and Chess, D. (2003). The Vision of Autonomic Computing. IEEE Computer, 36(1).",
    "[12] LeCun, Y., Bengio, Y., and Hinton, G. (2015). Deep Learning. Nature, 521(7553).",
    "[13] Chollet, F. (2018). Deep Learning with Python. Manning Publications.",
    "[14] Evans, E. (2003). Domain-Driven Design. Addison-Wesley Professional.",
    "[15] Gamma, E. et al. (1994). Design Patterns. Addison-Wesley.",
    "[16] Lundberg, S. and Lee, S. (2017). A Unified Approach to Interpreting Model Predictions. NeurIPS.",
    "[17] Devlin, J. et al. (2019). BERT: Pre-training of Deep Bidirectional Transformers. NAACL-HLT.",
    "[18] Silver, D. et al. (2016). Mastering the Game of Go. Nature, 529(7587).",
    "[19] Mnih, V. et al. (2015). Human-Level Control Through Deep RL. Nature, 518(7540).",
    "[20] Schulman, J. et al. (2017). Proximal Policy Optimization Algorithms. arXiv.",
]


# ============================================================
# MAIN GENERATION
# ============================================================
def main():
    print("Extracting content from Book 53...")
    chapters, chapter_order = extract_chapters(INPUT_FILE)
    print("Extracted %d chapters" % len(chapters))

    doc = Document()

    # ========================================
    # TITLE PAGE (raw, unformatted)
    # ========================================
    print("Creating title page...")
    add_raw_paragraph(doc, BOOK_TITLE, is_heading=True)
    add_raw_paragraph(doc, "")
    add_raw_paragraph(doc, AUTHOR_NAME)
    add_raw_paragraph(doc, "")
    add_raw_paragraph(doc, "Published by ECLEARNIX EDTECH PRIVATE LIMITED")
    add_raw_paragraph(doc, "Coimbatore, Tamil Nadu, India - 641049")
    add_raw_paragraph(doc, "")
    add_raw_paragraph(doc, "ISBN: 978-XX-XXXXXX-X-X")
    add_raw_paragraph(doc, "First Edition")
    add_raw_paragraph(doc, "Language: English")
    doc.add_page_break()

    # ========================================
    # TABLE OF CONTENTS (raw text placeholder)
    # ========================================
    print("Creating TOC...")
    add_raw_paragraph(doc, "Table of Contents", is_heading=True)
    for ch_idx, ch_title in enumerate(chapter_order, 1):
        add_raw_paragraph(doc, "%d. %s" % (ch_idx, ch_title))
    doc.add_page_break()

    # ========================================
    # LIST OF ABBREVIATIONS
    # ========================================
    print("Creating abbreviations...")
    add_raw_paragraph(doc, "List of Abbreviations", is_heading=True)
    add_raw_table(doc,
                  ["Abbreviation", "Full Form"],
                  [[a, f] for a, f in ABBREVIATIONS])
    add_raw_paragraph(doc, "")
    doc.add_page_break()

    # ========================================
    # CHAPTERS WITH CONTENT
    # ========================================
    for ch_idx, chapter_title in enumerate(chapter_order):
        chapter_num = ch_idx + 1
        print("Creating Chapter %d: %s..." % (chapter_num, chapter_title))

        # Chapter heading with "Chapter X" prefix (for rule engine detection)
        prefix_choice = random.choice([
            "CHAPTER %d\n" % chapter_num,
            "Chapter %d: " % chapter_num,
            "CHAPTER %d - " % chapter_num,
        ])
        add_raw_paragraph(doc, "%s%s" % (prefix_choice, chapter_title), is_heading=True)

        # Get chapter content
        paras = chapters.get(chapter_title, [])
        subheadings = SUBHEADINGS.get(chapter_num, [])
        chapter_tables = TABLES.get(chapter_num, [])
        figure_captions = FIGURE_CAPTIONS.get(chapter_num, [])
        bullets = BULLET_POINTS.get(chapter_num, [
            "This chapter explores critical concepts for modern enterprise AI",
            "Autonomous decision-making requires robust data foundations",
            "Integration with existing systems follows interoperability patterns",
            "Security and compliance must be addressed at every layer",
            "Human oversight ensures alignment with organizational goals",
            "Continuous learning is fundamental to self-evolving architectures",
            "Scalability requires both horizontal and vertical approaches",
            "The future lies in cognitive, self-managing platforms",
        ])

        # Distribute content across subheadings
        if paras and subheadings:
            paras_per_section = max(1, len(paras) // len(subheadings))
            table_idx = 0
            fig_idx = 0

            for sh_idx, subheading in enumerate(subheadings):
                # Subheading with section numbering (for rule engine: X.X pattern)
                add_raw_paragraph(doc, "%d.%d %s" % (chapter_num, sh_idx + 1, subheading), is_heading=True)

                # Get paragraphs for this section
                start = sh_idx * paras_per_section
                end = start + paras_per_section if sh_idx < len(subheadings) - 1 else len(paras)
                section_paras = paras[start:end]

                # Add body paragraphs
                for para_text in section_paras:
                    text = para_text.strip()
                    if text and (text[0].isdigit() and len(text) > 2 and text[1] in '.)'):
                        # Convert numbered items to bullet points
                        bullet_text = text
                        for pe in range(min(4, len(text))):
                            if text[pe] == ' ':
                                bullet_text = text[pe + 1:]
                                break
                        add_raw_paragraph(doc, "- %s" % bullet_text)
                    elif text.startswith('-') or text.startswith('•'):
                        add_raw_paragraph(doc, "- %s" % text.lstrip('-•').strip())
                    else:
                        add_raw_paragraph(doc, text)

                # Add table after certain sections
                if table_idx < len(chapter_tables) and sh_idx in [1, max(2, len(subheadings) // 2)]:
                    title, headers, rows = chapter_tables[table_idx]
                    add_raw_paragraph(doc, "Table %d.%d: %s" % (chapter_num, table_idx + 1, title))
                    add_raw_table(doc, headers, rows)
                    add_raw_paragraph(doc, "")
                    table_idx += 1

                # Add figure placeholder after certain sections
                if fig_idx < len(figure_captions) and sh_idx in [0, len(subheadings) // 3, 2 * len(subheadings) // 3]:
                    add_raw_paragraph(doc, "[Figure %d.%d: %s]" % (chapter_num, fig_idx + 1, figure_captions[fig_idx]))
                    add_raw_paragraph(doc, "Figure %d.%d: %s" % (chapter_num, fig_idx + 1, figure_captions[fig_idx]))
                    fig_idx += 1

            # Add remaining tables
            while table_idx < len(chapter_tables):
                title, headers, rows = chapter_tables[table_idx]
                add_raw_paragraph(doc, "Table %d.%d: %s" % (chapter_num, table_idx + 1, title))
                add_raw_table(doc, headers, rows)
                table_idx += 1

            # Add remaining figures
            while fig_idx < len(figure_captions):
                add_raw_paragraph(doc, "Figure %d.%d: %s" % (chapter_num, fig_idx + 1, figure_captions[fig_idx]))
                fig_idx += 1

        # KEY POINTS (bullet list)
        add_raw_paragraph(doc, "%d.%d Key Points" % (chapter_num, len(subheadings) + 1), is_heading=True)
        for bullet in bullets:
            add_raw_paragraph(doc, "- %s" % bullet)

        # MCQ BLOCK
        add_raw_paragraph(doc, "Multiple Choice Questions (MCQs)", is_heading=True)
        mcqs = generate_mcqs(chapter_num, chapter_title)
        for qi, mcq in enumerate(mcqs):
            q_num = qi + 1
            add_raw_paragraph(doc, "%d. %s" % (q_num, mcq['question']))
            opts = "\n".join(["a) %s" % mcq['options'][0],
                              "b) %s" % mcq['options'][1],
                              "c) %s" % mcq['options'][2],
                              "d) %s" % mcq['options'][3]])
            add_raw_paragraph(doc, opts)
            add_raw_paragraph(doc, "Answer: %s)" % mcq['answer'])

        doc.add_page_break()

    # ========================================
    # INDEX
    # ========================================
    print("Creating Index...")
    add_raw_paragraph(doc, "Index", is_heading=True)
    sorted_terms = sorted(INDEX_TERMS)
    current_letter = ""
    for term in sorted_terms:
        first = term[0].upper()
        if first != current_letter:
            current_letter = first
            add_raw_paragraph(doc, current_letter, is_heading=True)
        add_raw_paragraph(doc, term)

    doc.add_page_break()

    # ========================================
    # REFERENCES
    # ========================================
    print("Creating References...")
    add_raw_paragraph(doc, "References", is_heading=True)
    for ref in REFERENCES:
        add_raw_paragraph(doc, ref)

    # ========================================
    # SAVE
    # ========================================
    print("Saving raw unformatted document to %s..." % OUTPUT_FILE)
    doc.save(OUTPUT_FILE)

    total_paras = len(doc.paragraphs)
    total_words = sum(len(p.text.split()) for p in doc.paragraphs if p.text.strip())
    print("\n=== RAW DOCUMENT STATISTICS ===")
    print("Total paragraphs: %d" % total_paras)
    print("Total words: %d" % total_words)
    print("Estimated pages (unformatted): ~%d" % (total_words // 250))
    print("Tables: %d" % len(doc.tables))
    print("\nRaw unformatted document saved!")
    print("Run the formatter: python run.py sample_manuscript.docx")


if __name__ == "__main__":
    main()
