import os
import io
import random
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

random.seed(42)

FONTS = ['Arial', 'Calibri', 'Times New Roman', 'Courier New', 'Verdana', 'Georgia']
SIZES = [Pt(9), Pt(10), Pt(11), Pt(12), Pt(13), Pt(14), Pt(16)]
ALIGNMENTS = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY]
SPACINGS = [1.0, 1.15, 1.5, 2.0]

ACADEMIC_SENTENCES = [
    "Self-evolving enterprise architectures represent a paradigm shift from static, brittle infrastructures to dynamic, autonomous systems capable of continuous self-adaptation.",
    "The integration of reinforcement learning with distributed event meshes enables real-time optimization of computational workloads across multi-cloud environments.",
    "Traditional perimeter-based security models fail to address modern lateral threat vectors in highly decentralized, service-oriented platforms.",
    "Cognitive resource scheduling dynamically reallocates containerized microservices based on predictive telemetry and multivariate anomaly detection.",
    "By leveraging generative AI for automated schema migration and semantic data reconciliation, enterprises eliminate substantial technical debt.",
    "The biological metaphor of homeostatic regulation provides a rigorous mathematical framework for designing fault-tolerant, self-healing software clusters.",
    "Decentralized data fabrics unify heterogeneous relational and graph databases into a unified, high-throughput semantic querying interface.",
    "Zero-trust security architectures enforce continuous cryptographic verification of all internal service requests regardless of network topology.",
    "Empirical evaluations across high-frequency financial trading benchmarks demonstrate a 42% reduction in p99 latency using cognitive edge caching.",
    "Explainable artificial intelligence (XAI) frameworks provide human operators with verifiable decision paths for autonomous remediation actions.",
    "A comprehensive analysis of distributed consensus mechanisms reveals critical scalability bottlenecks under Byzantine fault conditions.",
    "Continuous feedback loops governed by policy gradient algorithms dynamically adjust rate limits and firewall thresholds in response to evolving threat signatures.",
    "Microservice deconstruction demands sophisticated service meshes equipped with automated mutual TLS (mTLS) certificate rotation.",
    "The data lifecycle encompasses secure edge ingestion, stream normalization, contextual enrichment, and privacy-preserving federated aggregation.",
    "DevSecOps pipelines must integrate automated formal verification proofs prior to deploying self-modifying microservice code to production environments."
]

def create_sample_image(title_text: str) -> io.BytesIO:
    """Generates an architectural block diagram image in memory."""
    img = Image.new('RGB', (600, 260), color=(240, 245, 250))
    draw = ImageDraw.Draw(img)
    
    # Outer border
    draw.rectangle([5, 5, 595, 255], outline=(50, 80, 120), width=3)
    
    # Diagram Boxes
    boxes = [
        (30, 40, 180, 120, "Edge Gateway\nIngestion Layer", (200, 220, 240)),
        (220, 40, 380, 120, "Cognitive AI Core\nInference & Policy", (180, 230, 200)),
        (420, 40, 570, 120, "Distributed Mesh\nData Fabric", (240, 220, 200)),
        (120, 150, 480, 230, "Autonomous Feedback Loop & Self-Healing Controller", (230, 210, 240))
    ]
    
    for x1, y1, x2, y2, label, bg in boxes:
        draw.rectangle([x1, y1, x2, y2], fill=bg, outline=(40, 60, 90), width=2)
        draw.text((x1 + 10, y1 + 25), label, fill=(10, 20, 40))
        
    # Arrows / Lines
    draw.line([180, 80, 220, 80], fill=(20, 40, 80), width=3)
    draw.line([380, 80, 420, 80], fill=(20, 40, 80), width=3)
    draw.line([300, 120, 300, 150], fill=(20, 40, 80), width=3)
    
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return buf

def get_random_sentence():
    return random.choice(ACADEMIC_SENTENCES)

def get_rich_paragraph(num_sentences=None):
    if num_sentences is None:
        num_sentences = random.randint(8, 16)
    return " ".join([get_random_sentence() for _ in range(num_sentences)])

def apply_random_formatting(paragraph, is_heading=False):
    font_name = random.choice(FONTS)
    font_size = random.choice(SIZES)
    alignment = random.choice(ALIGNMENTS)
    spacing = random.choice(SPACINGS)
    
    paragraph.alignment = alignment
    if paragraph.paragraph_format:
        paragraph.paragraph_format.line_spacing = spacing
        if not is_heading and random.random() > 0.5:
            paragraph.paragraph_format.first_line_indent = None
        elif is_heading:
            paragraph.paragraph_format.first_line_indent = None
            
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = font_size
        if is_heading:
            run.bold = random.choice([True, False])
            if random.random() > 0.6:
                run.text = run.text.upper()
        else:
            run.bold = False

def add_p(doc, text, is_heading=False):
    p = doc.add_paragraph()
    p.add_run(text)
    apply_random_formatting(p, is_heading=is_heading)
    return p

def generate_sample(output_path):
    doc = Document()
    
    # ---------------- PAGE 1: TITLE & AUTHOR PAGE ----------------
    add_p(doc, "Self-Evolving Enterprise Systems: AI-Driven Architectures for the Next Generation of Digital Platforms", is_heading=True)
    add_p(doc, "Vinod Kumar Thottempudi", is_heading=False)
    add_p(doc, "Site Reliability Engineer", is_heading=False)
    add_p(doc, "Published by", is_heading=False)
    add_p(doc, "ECLEARNIX EDTECH PRIVATE LIMITED", is_heading=False)
    add_p(doc, "Coimbatore, Tamil Nadu, India – 641049", is_heading=False)
    doc.add_page_break()
    
    # ---------------- PAGE 2: PREFACE & ABSTRACT ----------------
    add_p(doc, "PREFACE", is_heading=True)
    add_p(doc, get_rich_paragraph(14))
    add_p(doc, get_rich_paragraph(12))
    add_p(doc, "ABOUT THE AUTHOR", is_heading=True)
    add_p(doc, get_rich_paragraph(10))
    doc.add_page_break()
    
    # ---------------- TABLE OF CONTENTS / INDEX ----------------
    add_p(doc, "INDEX", is_heading=True)
    add_p(doc, "Preface\txiii", is_heading=False)
    add_p(doc, "About the Author\txvi", is_heading=False)
    add_p(doc, "List of Abbreviations\txvii", is_heading=False)
    
    chapters_metadata = [
        "The Dawn of Self-Evolving Systems",
        "From Monoliths to Intelligent Mesh",
        "Technological Foundations of AI-Driven Platforms",
        "Cognitive Resource Management and Optimization",
        "Engineering Resilience: Self-Healing and Autonomy",
        "Generative Architectures: AI-Assisted Schema and Code",
        "The Intelligent Data Fabric",
        "Reinforcement Learning for Continuous Improvement",
        "Securing the Autonomous Perimeter",
        "Explainable AI and Deep Observability",
        "Scaling the Self-Evolving Enterprise",
        "AIOps: Reimagining the Software Lifecycle",
        "Governance, Ethics, and Human Agency",
        "Real-World Implementation and Case Studies",
        "The Horizon of Enterprise Intelligence",
        "High-Throughput Distributed Consensus",
        "Zero-Trust Edge Gateways and Policy Engines",
        "Predictive Failure Analysis in Cloud Infrastructure",
        "Quantum-Safe Enterprise Cryptography",
        "Synthesis, Future Paradigms, and Roadmap"
    ]
    
    page_counter = 1
    for ch_num, ch_title in enumerate(chapters_metadata, 1):
        add_p(doc, f"CHAPTER - {ch_num}\t{page_counter}", is_heading=False)
        add_p(doc, f"{ch_title}\t{page_counter}", is_heading=False)
        for sec_i in range(1, 4):
            add_p(doc, f"{ch_num}.{sec_i}\tTheoretical Framework and Analysis\t{page_counter + sec_i * 3}", is_heading=False)
        add_p(doc, f"Multiple Choice Questions (MCQs)\t{page_counter + 18}", is_heading=False)
        page_counter += 24
    doc.add_page_break()
    
    # ---------------- 20 EXTENSIVE CHAPTERS (500-PAGE VOLUME) ----------------
    for ch_idx, ch_title in enumerate(chapters_metadata, 1):
        # Chapter Heading in messy unformatted style
        chap_header = random.choice([
            f"CHAPTER - {ch_idx}\n{ch_title}",
            f"CHAPTER {ch_idx}: {ch_title}",
            f"{ch_idx}. {ch_title.upper()}",
            f"Chapter {ch_idx} - {ch_title}"
        ])
        add_p(doc, chap_header, is_heading=True)
        
        # Chapter Intro
        add_p(doc, "Introduction", is_heading=True)
        for _ in range(random.randint(3, 5)):
            add_p(doc, get_rich_paragraph(random.randint(10, 16)))
            
        # 4 to 6 Sections per chapter
        num_sections = random.randint(4, 6)
        for sec_idx in range(1, num_sections + 1):
            sec_name = f"{ch_idx}.{sec_idx} Architectural Principles of {get_random_sentence().split()[1]} {get_random_sentence().split()[2]}"
            add_p(doc, sec_name, is_heading=True)
            
            # Section body paragraphs
            for _ in range(random.randint(4, 6)):
                add_p(doc, get_rich_paragraph(random.randint(12, 18)))
                
            # 2 Subsections per section
            for sub_idx in range(1, 3):
                subsec_name = f"{ch_idx}.{sec_idx}.{sub_idx} Methodological Implementation and Validation"
                add_p(doc, subsec_name, is_heading=True)
                
                for _ in range(random.randint(4, 7)):
                    add_p(doc, get_rich_paragraph(random.randint(12, 20)))
                    
                # Hyphen-based list to test bullet conversion
                if random.random() > 0.3:
                    for l_idx in range(1, random.randint(3, 6)):
                        add_p(doc, f"- Core operational constraint {l_idx}: {get_random_sentence()}")
                        
            # Add Table with unformatted caption
            if random.random() > 0.2:
                tbl_title = f"Table {ch_idx}.{sec_idx}: Performance Evaluation Matrix for Variant {sec_idx}"
                p_tbl = doc.add_paragraph()
                p_tbl.add_run(tbl_title)
                apply_random_formatting(p_tbl, is_heading=False)
                
                rows_cnt = random.randint(4, 6)
                cols_cnt = random.randint(3, 4)
                tbl = doc.add_table(rows=rows_cnt, cols=cols_cnt)
                headers = ["Subsystem Component", "Throughput (RPS)", "P99 Latency (ms)", "CPU Utilization (%)"]
                for c_i in range(cols_cnt):
                    tbl.cell(0, c_i).text = headers[c_i % len(headers)]
                for r_i in range(1, rows_cnt):
                    for c_i in range(cols_cnt):
                        tbl.cell(r_i, c_i).text = f"Metric_{r_i}_{c_i}_{random.randint(80, 99)}.{random.randint(10, 99)}"
                doc.add_paragraph() # space
                
            # Add Image Figure with caption
            if random.random() > 0.2:
                img_stream = create_sample_image(f"Chapter {ch_idx} Architecture")
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_img = p_img.add_run()
                run_img.add_picture(img_stream, width=Inches(4.8))
                
                fig_cap = f"Figure {ch_idx}.{sec_idx}: System architecture diagram illustrating cognitive feedback mechanisms."
                p_cap = doc.add_paragraph()
                p_cap.add_run(fig_cap)
                apply_random_formatting(p_cap, is_heading=False)
                
        # ---------------- MCQ BLOCK AT END OF EVERY CHAPTER ----------------
        add_p(doc, "Multiple Choice Questions (MCQs)", is_heading=True)
        for q_num in range(1, random.randint(6, 10)):
            # Question line
            q_text = f"{q_num}. Which of the following best describes the core design goal of {ch_title} in enterprise digital platforms?"
            p_q = doc.add_paragraph()
            p_q.add_run(q_text)
            apply_random_formatting(p_q, is_heading=True)
            
            # Options (multiline or separate)
            options_text = (
                f"a) Implementing rigid monolithic perimeter controls\n"
                f"b) Enabling autonomous real-time self-evolution and decentralized optimization\n"
                f"c) Eliminating all observability telemetry across services\n"
                f"d) Restricting deployment pipelines to manual physical interventions"
            )
            p_opt = doc.add_paragraph()
            p_opt.add_run(options_text)
            apply_random_formatting(p_opt, is_heading=False)
            
            # Answer
            ans_text = f"Answer: {random.choice(['a', 'b', 'c', 'd'])})"
            p_ans = doc.add_paragraph()
            p_ans.add_run(ans_text)
            apply_random_formatting(p_ans, is_heading=False)
            
        doc.add_page_break()

    # ---------------- COMPREHENSIVE REFERENCES SECTION ----------------
    add_p(doc, "References", is_heading=True)
    for ref_i in range(1, 130):
        year = random.randint(2014, 2026)
        vol = random.randint(1, 55)
        pp1 = random.randint(100, 800)
        pp2 = pp1 + random.randint(8, 28)
        if random.random() > 0.5:
            r_text = f"[{ref_i}] T. K. Author{ref_i}, M. L. Coauthor{ref_i}, and S. R. Researcher{ref_i}, \"Advanced architectures for cognitive enterprise systems and autonomous orchestration,\" IEEE Transactions on Cloud Computing, vol. {vol}, no. 2, pp. {pp1}-{pp2}, {year}."
        else:
            r_text = f"Author{ref_i}, T. K., Coauthor{ref_i}, M. L., & Researcher{ref_i}, S. R. ({year}). Design paradigms for self-healing distributed digital platforms. Journal of Systems and Software Architecture, {vol}(3), {pp1}-{pp2}. https://doi.org/10.1016/j.sysarc.{year}.{ref_i:04d}"
        add_p(doc, r_text, is_heading=False)
        
    doc.add_page_break()
    
    # ---------------- SUBJECT INDEX AT THE END ----------------
    add_p(doc, "Subject Index", is_heading=True)
    index_terms = [
        "AIOps", "Anomaly Detection", "Autonomous Remediation", "Byzantine Fault Tolerance",
        "Cloud Native", "Cognitive Mesh", "Data Fabric", "DevSecOps", "Distributed Consensus",
        "Edge Computing", "Explainable AI", "Federated Learning", "Generative Schema",
        "Microservices", "Neural Routing", "Observability", "Policy Engines", "Reinforcement Learning",
        "Self-Healing Infrastructure", "Telemetry Pipelines", "Zero Trust Architecture"
    ]
    for term in index_terms:
        pages = ", ".join([str(random.randint(10, 480)) for _ in range(random.randint(3, 7))])
        add_p(doc, f"{term}\t{pages}", is_heading=False)
        
    doc.save(output_path)

if __name__ == "__main__":
    output_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'sample_manuscript.docx')
    print(f"Generating 500-page unformatted manuscript at {output_path}...")
    generate_sample(output_path)
    print(f"Generated successfully at {output_path}!")

