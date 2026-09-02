from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from src.config import ElementType, FormattingConfig, DEFAULT_CONFIG
import re

def add_page_number_to_run(run):
    """Inserts a dynamic Word PAGE field into a run."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

class DocumentFormatter:
    def __init__(self, config: FormattingConfig = DEFAULT_CONFIG):
        self.config = config

    def setup_document(self, doc: Document, book_title: str = "", author_name: str = "") -> None:
        for section in doc.sections:
            # Layout & Margins: Top: 1.52 cm, Bottom: 1.52 cm, Left: 1.97 cm, Right: 1.96 cm
            section.top_margin = Cm(1.52)
            section.bottom_margin = Cm(1.52)
            section.left_margin = Cm(1.97)
            section.right_margin = Cm(1.96)
            section.gutter = Cm(0)
            
            # Dual-sided footer: Left = Author/Book Title, Right = Page Number
            self._setup_footer(section, book_title, author_name)

    def _setup_footer(self, section, book_title: str = "", author_name: str = "") -> None:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = "" # clear existing text
        
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(6)
        pf.space_after = Pt(0)
        
        # Left side: Author Title & Book Title
        left_text = ""
        if author_name and book_title:
            left_text = f"{author_name} | {book_title}"
        elif book_title:
            left_text = book_title
        elif author_name:
            left_text = author_name
        else:
            left_text = "Publication Manuscript"
            
        if len(left_text) > 65:
            left_text = left_text[:62] + "..."
            
        r_title = p.add_run(left_text)
        r_title.font.name = self.config.font_name
        r_title.font.size = Pt(9)
        r_title.font.color.rgb = RGBColor(80, 80, 80)
        
        # Right tab to right margin
        r_tab = p.add_run("\t")
        
        # Right side: Dynamic Page Number
        r_page = p.add_run()
        r_page.font.name = self.config.font_name
        r_page.font.size = Pt(10)
        r_page.font.bold = True
        r_page.font.color.rgb = RGBColor(0, 0, 0)
        add_page_number_to_run(r_page)
        
        # Add a right-aligned tab stop at 9678 dxa (~17.07 cm)
        pPr = p._p.get_or_add_pPr()
        tabs = parse_xml(r'<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:tab w:val="right" w:pos="9678"/></w:tabs>')
        pPr.append(tabs)

    def format_paragraph(self, paragraph, element_type: ElementType) -> None:
        pf = paragraph.paragraph_format
        
        # Defaults
        alignment = WD_ALIGN_PARAGRAPH.LEFT
        space_before = Pt(0)
        space_after = Pt(0)
        line_spacing = 1.5  # STRICT 1.5 Line Spacing default across body and lists
        first_line_indent = None
        left_indent = None
        
        font_size = Pt(12)
        bold = False
        italic = False
        text_color = RGBColor(0, 0, 0)
        heading_blue = RGBColor(31, 78, 120)  # Deep Academic Blue
        
        if element_type == ElementType.TITLE:
            font_size = Pt(24)
            bold = True
            alignment = WD_ALIGN_PARAGRAPH.CENTER
            space_before = Pt(48)
            space_after = Pt(140)
            line_spacing = 1.15
            
        elif element_type == ElementType.AUTHOR:
            raw_text = getattr(paragraph, 'text', '').strip()
            line_spacing = 1.15
            if re.match(r'^\s*Published\s+by\s*$', raw_text, re.I):
                font_size = Pt(11)
                bold = True
                alignment = WD_ALIGN_PARAGRAPH.CENTER
                space_before = Pt(0)
                space_after = Pt(6)
            elif re.search(r'ECLEARNIX|EDTECH|PRIVATE LIMITED|PUBLISH|PRESS|PUBLICATIONS', raw_text, re.I):
                font_size = Pt(12)
                bold = True
                alignment = WD_ALIGN_PARAGRAPH.CENTER
                space_before = Pt(0)
                space_after = Pt(6)
            elif re.search(r'Coimbatore|Tamil\s*Nadu|India|DOI:|ISBN|Postal', raw_text, re.I):
                font_size = Pt(10)
                bold = False
                alignment = WD_ALIGN_PARAGRAPH.CENTER
                space_before = Pt(0)
                space_after = Pt(0)
            elif re.search(r'Engineer|Architect|Professor|Director|Lead|Scientist|Department|University|Laboratory', raw_text, re.I):
                font_size = Pt(12)
                bold = False
                alignment = WD_ALIGN_PARAGRAPH.CENTER
                space_before = Pt(0)
                space_after = Pt(160)
            else:
                font_size = Pt(20)
                bold = True
                alignment = WD_ALIGN_PARAGRAPH.CENTER
                space_before = Pt(0)
                space_after = Pt(12)
                
        elif element_type == ElementType.CHAPTER_HEADING:
            font_size = Pt(16)
            bold = True
            alignment = WD_ALIGN_PARAGRAPH.LEFT
            space_before = Pt(36)
            space_after = Pt(18)
            line_spacing = 1.15
            text_color = heading_blue
            pf.page_break_before = True
            
        elif element_type in (ElementType.SECTION_HEADING, ElementType.SUBSECTION_HEADING):
            font_size = Pt(14) if element_type == ElementType.SECTION_HEADING else Pt(12)
            bold = True
            alignment = WD_ALIGN_PARAGRAPH.LEFT
            space_before = Pt(18)
            space_after = Pt(6)
            line_spacing = 1.15
            text_color = heading_blue
            
        elif element_type in (ElementType.REFERENCE_HEADING, ElementType.TOC_HEADING, 
                              ElementType.INDEX_HEADING, ElementType.ABSTRACT_HEADING, 
                              ElementType.ABOUT_AUTHOR_HEADING):
            # Special Unique Titles & MCQ Title: CENTER ALIGNED
            font_size = Pt(16)
            bold = True
            alignment = WD_ALIGN_PARAGRAPH.CENTER
            space_before = Pt(36)
            space_after = Pt(18)
            line_spacing = 1.15
            text_color = RGBColor(0, 0, 0)
            pf.page_break_before = True
            
        elif element_type == ElementType.MCQ_HEADING:
            # MCQ Title: CENTER ALIGNED
            font_size = Pt(16)
            bold = True
            alignment = WD_ALIGN_PARAGRAPH.CENTER
            space_before = Pt(36)
            space_after = Pt(18)
            line_spacing = 1.15
            text_color = heading_blue
            pf.page_break_before = True
            
        elif element_type == ElementType.BODY:
            # Body Paragraphs: 12 pt, Justified, STRICT 1.5 line spacing, 1.27 cm indent
            font_size = Pt(12)
            alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            line_spacing = 1.5
            first_line_indent = Cm(1.27)
            space_after = Pt(6)
            
        elif element_type == ElementType.LIST_ITEM:
            # List Items: 12 pt, STRICT 1.5 line spacing, Left indent 1.27 cm
            font_size = Pt(12)
            alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            left_indent = Cm(1.27)
            first_line_indent = -Cm(0.5)
            line_spacing = 1.5
            space_after = Pt(4)
            
        elif element_type in (ElementType.FIGURE_CAPTION, ElementType.TABLE_CAPTION):
            font_size = Pt(10)
            italic = True
            alignment = WD_ALIGN_PARAGRAPH.CENTER
            space_before = Pt(6)
            space_after = Pt(12)
            line_spacing = 1.15
            
        elif element_type == ElementType.REFERENCE_ENTRY:
            font_size = Pt(10)
            alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            first_line_indent = -Cm(0.75)
            left_indent = Cm(0.75)
            line_spacing = 1.5
            space_after = Pt(4)
            
        elif element_type in (ElementType.TOC_ENTRY, ElementType.INDEX_ENTRY):
            font_size = Pt(11)
            alignment = WD_ALIGN_PARAGRAPH.LEFT
            line_spacing = 1.15
            space_after = Pt(3)
            
        elif element_type == ElementType.MCQ_QUESTION:
            # MCQ Question: Normal Text (Not Bold, Sentence Case), 1.5 line spacing
            font_size = Pt(12)
            bold = False
            alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            space_before = Pt(12)
            space_after = Pt(4)
            line_spacing = 1.5
            
        elif element_type == ElementType.MCQ_OPTION:
            # MCQ Option: Normal Text with Capital Option Letter (A, B, C, D), 1.5 line spacing
            font_size = Pt(12)
            bold = False
            alignment = WD_ALIGN_PARAGRAPH.LEFT
            left_indent = Cm(0.75)
            line_spacing = 1.5
            space_after = Pt(4)
            
        elif element_type == ElementType.MCQ_ANSWER:
            font_size = Pt(12)
            bold = True
            alignment = WD_ALIGN_PARAGRAPH.LEFT
            left_indent = Cm(0.75)
            space_before = Pt(2)
            space_after = Pt(12)
            line_spacing = 1.5
            
        elif element_type == ElementType.BLANK:
            space_after = Pt(0)
            
        pf.alignment = alignment
        pf.space_before = space_before
        pf.space_after = space_after
        pf.line_spacing = line_spacing
        if first_line_indent is not None:
            pf.first_line_indent = first_line_indent
        if left_indent is not None:
            pf.left_indent = left_indent
            
        for run in paragraph.runs:
            run.font.name = self.config.font_name
            run.font.size = font_size
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = text_color

    def format_table(self, table) -> None:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.name = self.config.font_name
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        if i == 0:
                            run.font.bold = True
