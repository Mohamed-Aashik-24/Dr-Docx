import io
import re
from typing import List
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from src.config import ElementType, DEFAULT_CONFIG, FormattingConfig
from src.parser import ParsedElement, ParsedTable
from src.structure_recognizer import ClassifiedElement
from src.formatter import DocumentFormatter

def normalize_question_text(text: str) -> str:
    """Converts ALL CAPS MCQ questions into Normal Sentence Case while retaining acronyms."""
    stripped = text.strip()
    if not stripped:
        return text
        
    m = re.match(r'^(\s*(?:Q?\d+[\.\)]\s*)?)(.*)$', stripped, re.DOTALL)
    if not m:
        return text
        
    prefix = m.group(1)
    body = m.group(2).strip()
    
    if not body:
        return text
        
    if body.isupper():
        body_sc = body[0].upper() + body[1:].lower()
        acronyms = ['AI', 'AIOPS', 'CDN', 'WAF', 'IP', 'API', 'LLM', 'CPU', 'RAM', 'TCP', 'UDP', 'DNS', 'URL', 'HTTPS', 'HTTP', 'JWT', 'SIEM', 'IoT', 'ZTNA']
        for ac in acronyms:
            body_sc = re.sub(r'\b' + re.escape(ac.lower()) + r'\b', ac, body_sc, flags=re.IGNORECASE)
        return prefix + body_sc
        
    return text

def normalize_option_text(text: str) -> str:
    """Ensures MCQ option letters are CAPITALized (A), B), C), D)) across single or multiline strings."""
    if not text:
        return text
        
    def replace_prefix(m):
        prefix_space = m.group(1) or ''
        letter = m.group(2).upper()
        closing = m.group(3) or ')'
        return f"{prefix_space}{letter}{closing} "

    # Replace all option letters at line start or after newline (e.g. a) -> A), \nb) -> \nB))
    res = re.sub(r'(^|\n)\s*([a-dA-D])([\)\.])\s*', replace_prefix, text)
    return res

class DocxGenerator:
    def __init__(self, config: FormattingConfig = DEFAULT_CONFIG):
        self.formatter = DocumentFormatter(config)

    def generate(self, classified_elements: List[ClassifiedElement], output_path: str) -> str:
        doc = Document()
        
        # 1. Extract book title and author name for front page and footers
        book_title = "Publication Manuscript"
        author_name = ""
        
        for ce in classified_elements:
            if isinstance(ce.element, ParsedTable):
                continue
            text = getattr(ce.element, 'text', '').strip()
            if not text:
                continue
            if ce.element_type == ElementType.TITLE and book_title == "Publication Manuscript":
                book_title = text
            elif ce.element_type == ElementType.AUTHOR and not author_name:
                if not re.search(r'Published by|ECLEARNIX|EDTECH|Coimbatore|Tamil|DOI:|ISBN|Engineer|Professor|Department', text, re.I):
                    author_name = text
                    
        self.formatter.setup_document(doc, book_title=book_title, author_name=author_name)
        
        current_chapter = 1
        chapter_table_idx = 0
        chapter_fig_idx = 0
        
        title_page_active = False
        title_page_completed = False
        i = 0
        consecutive_blank_count = 0
        
        while i < len(classified_elements):
            ce = classified_elements[i]
            element = ce.element
            element_type = ce.element_type
            
            # Track Title Page state
            if element_type in (ElementType.TITLE, ElementType.AUTHOR):
                title_page_active = True
            elif title_page_active and not title_page_completed:
                doc.add_page_break()
                title_page_active = False
                title_page_completed = True

            # Page Break before About the Author section
            if element_type == ElementType.ABOUT_AUTHOR_HEADING and i > 0 and title_page_completed:
                doc.add_page_break()

            # Check for TOC or Index block
            if element_type in (ElementType.TOC_ENTRY, ElementType.INDEX_ENTRY):
                toc_rows = []
                while i < len(classified_elements) and classified_elements[i].element_type in (ElementType.TOC_ENTRY, ElementType.INDEX_ENTRY):
                    entry_text = getattr(classified_elements[i].element, 'text', '').strip()
                    if entry_text:
                        if '\t' in entry_text:
                            parts = entry_text.rsplit('\t', 1)
                            toc_rows.append((parts[0].strip(), parts[1].strip()))
                        else:
                            toc_rows.append((entry_text, ""))
                    i += 1
                
                # Create neat 2-column table
                if toc_rows:
                    tbl = doc.add_table(rows=len(toc_rows), cols=2)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for row_idx, (col1_txt, col2_txt) in enumerate(toc_rows):
                        row = tbl.rows[row_idx]
                        
                        c0 = row.cells[0]
                        c0.width = Inches(5.6)
                        p0 = c0.paragraphs[0]
                        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p0.paragraph_format.space_before = Pt(2)
                        p0.paragraph_format.space_after = Pt(2)
                        p0.paragraph_format.line_spacing = 1.0
                        r0 = p0.add_run(col1_txt)
                        r0.font.name = self.formatter.config.font_name
                        r0.font.size = Pt(11)
                        if col1_txt.startswith('CHAPTER') or col1_txt.startswith('Chapter') or col1_txt in ('Preface', 'About the Author', 'List of Abbreviations'):
                            r0.font.bold = True
                            
                        c1 = row.cells[1]
                        c1.width = Inches(1.1)
                        p1 = c1.paragraphs[0]
                        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        p1.paragraph_format.space_before = Pt(2)
                        p1.paragraph_format.space_after = Pt(2)
                        p1.paragraph_format.line_spacing = 1.0
                        r1 = p1.add_run(col2_txt)
                        r1.font.name = self.formatter.config.font_name
                        r1.font.size = Pt(11)
                        if col1_txt.startswith('CHAPTER') or col1_txt.startswith('Chapter'):
                            r1.font.bold = True
                    
                    p_space = doc.add_paragraph()
                    p_space.paragraph_format.space_before = Pt(0)
                    p_space.paragraph_format.space_after = Pt(12)
                continue
                
            # Track chapter index
            if element_type == ElementType.CHAPTER_HEADING:
                raw_text = getattr(element, 'text', '')
                m = re.search(r'\b(?:chapter|ch\.?)\s*[-:]?\s*(\d+)', raw_text, re.I)
                if m:
                    current_chapter = int(m.group(1))
                else:
                    current_chapter += 1
                chapter_table_idx = 0
                chapter_fig_idx = 0
            
            if isinstance(element, ParsedTable):
                table = doc.add_table(rows=element.num_rows, cols=element.num_cols)
                for r_idx, row in enumerate(element.rows):
                    for j, cell_text in enumerate(row):
                        table.cell(r_idx, j).text = cell_text
                self.formatter.format_table(table)
                p_trail = doc.add_paragraph()
                p_trail.paragraph_format.space_before = Pt(0)
                p_trail.paragraph_format.space_after = Pt(6)
            
            elif isinstance(element, ParsedElement):
                # Reset blank counter when we encounter non-blank content
                if element_type != ElementType.BLANK:
                    consecutive_blank_count = 0

                if element_type == ElementType.BLANK:
                    consecutive_blank_count += 1
                    if consecutive_blank_count <= 1:
                        # Allow at most one blank paragraph between content
                        paragraph = doc.add_paragraph()
                        self.formatter.format_paragraph(paragraph, element_type)
                    i += 1
                    continue
                elif element.has_image and element.image_data:
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    image_stream = io.BytesIO(element.image_data)
                    run = paragraph.add_run()
                    try:
                        run.add_picture(image_stream, width=Inches(4.8))
                    except Exception:
                        pass
                else:
                    text_content = element.text
                    
                    # Convert hyphens in list items to bullet points (•)
                    if element_type == ElementType.LIST_ITEM:
                        text_content = re.sub(r'^\s*[-*–▪]\s*', '• ', text_content)
                        if not text_content.startswith('• ') and not re.match(r'^\s*(\d+\.|[a-zA-Z]\))', text_content):
                            text_content = '• ' + text_content.lstrip()
                            
                    # Normalize MCQ Questions (Normal Sentence Case, not ALL CAPS)
                    elif element_type == ElementType.MCQ_QUESTION:
                        text_content = normalize_question_text(text_content)
                        
                    # Normalize MCQ Options (Capital Letters A), B), C), D))
                    elif element_type == ElementType.MCQ_OPTION:
                        text_content = normalize_option_text(text_content)
                            
                    # Auto-number Table Captions
                    elif element_type == ElementType.TABLE_CAPTION:
                        chapter_table_idx += 1
                        clean_caption = re.sub(r'^\s*(Table|Tbl\.?)\s*(\d+(\.\d+)?)?[\s:-]*', '', text_content, flags=re.I).strip()
                        text_content = f"Table {current_chapter}.{chapter_table_idx}: {clean_caption}"
                        
                    # Auto-number Figure Captions
                    elif element_type == ElementType.FIGURE_CAPTION:
                        chapter_fig_idx += 1
                        clean_caption = re.sub(r'^\s*(Figure|Fig\.?)\s*(\d+(\.\d+)?)?[\s:-]*', '', text_content, flags=re.I).strip()
                        text_content = f"Figure {current_chapter}.{chapter_fig_idx}: {clean_caption}"
                        
                    paragraph = doc.add_paragraph()
                    paragraph.text = text_content
                    self.formatter.format_paragraph(paragraph, element_type)
                    
            i += 1
                    
        doc.save(output_path)
        return output_path

    def _add_page_break(self, doc) -> None:
        doc.add_page_break()
